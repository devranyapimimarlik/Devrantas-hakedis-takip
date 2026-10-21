from io import BytesIO
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage

LOGO = Path(__file__).parent / "fonts" / "logo.png"
COMPANY = {"name": "DEVRANTAŞ YAPI MİMARLIK İNŞAAT LTD. ŞTİ.", "address": "Nurtanesi Sok. No: 3/1, Merkez Mah., Çekmeköy, İstanbul",
           "phone": "Tel: 0533 482 57 02", "email": "E-posta: Devrantasyapimimarlik@gmail.com"}
GOLD = "B8860B"
THIN = Side(style="thin", color="E2D3B4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '#,##0.00 "₺"'


def safe_title(t: str) -> str:
    return "".join("-" if ch in '\\/*?:[]' else ch for ch in t)[:28]


def _w(ws, row, col, val, bold=False, size=10, color="1C1917", fill=None, align=None, fmt=None, border=False):
    c = ws.cell(row=row, column=col, value=val)
    c.font = Font(bold=bold, size=size, color=color, name="Arial")
    if fill: c.fill = PatternFill("solid", fgColor=fill)
    if align: c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    if fmt: c.number_format = fmt
    if border: c.border = BORDER
    return c


async def refresh_company():
    from routers.settings import company_info, logo_path
    COMPANY.update(await company_info())
    COMPANY["logo"] = await logo_path()


def build_table(title: str, subtitle: str, columns: list, rows: list, summary: list = None) -> bytes:
    wb = Workbook(); ws = wb.active; ws.title = safe_title(title)
    widths = [max(14, min(50, len(str(c)) + 6)) for c in columns]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w if i > 1 else max(w, 30)
    img = XLImage(COMPANY.get("logo") or str(LOGO)); img.width = img.height = 70
    ws.add_image(img, "A1")
    for rr in (1, 2, 3, 4): ws.row_dimensions[rr].height = 16 if rr > 1 else 28
    _w(ws, 1, 2, COMPANY.get("upper") or COMPANY["name"], bold=True, size=14, color=GOLD)
    _w(ws, 2, 2, COMPANY["address"], size=9, color="57534E"); _w(ws, 3, 2, COMPANY["phone"], size=9, color="57534E"); _w(ws, 4, 2, COMPANY["email"], size=9, color="57534E")
    _w(ws, 6, 1, title, bold=True, size=16, color=GOLD); _w(ws, 7, 1, subtitle, size=9, color="57534E")
    r = 9
    if summary:
        for k, v in summary:
            _w(ws, r, 1, k, bold=True, color=GOLD, size=9); _w(ws, r, 2, v, fmt=MONEY if isinstance(v, (int, float)) else None, align="left"); r += 1
        r += 1
    for i, h in enumerate(columns, 1):
        _w(ws, r, i, h, bold=True, fill="FAF5E8", border=True)
    r += 1
    for row in rows:
        for i, v in enumerate(row, 1):
            _w(ws, r, i, v, border=True, fmt=MONEY if isinstance(v, (int, float)) and not isinstance(v, bool) else None, align="right" if isinstance(v, (int, float)) else "left")
        r += 1
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True
    buf = BytesIO(); wb.save(buf)
    return buf.getvalue()


def build_document(title: str, meta: list, items: list, totals: list, sections: list, signatures=None) -> bytes:
    signatures = signatures or ("Müşteri Onayı", f"{COMPANY.get('upper', '')}\nKaşe / İmza")
    wb = Workbook(); ws = wb.active; ws.title = safe_title(title)
    for col, w in zip("ABCDEF", [6, 48, 12, 14, 18, 20]):
        ws.column_dimensions[col].width = w
    img = XLImage(COMPANY.get("logo") or str(LOGO)); img.width = img.height = 70
    ws.add_image(img, "A1")
    ws.row_dimensions[1].height = 28; ws.row_dimensions[2].height = 16; ws.row_dimensions[3].height = 16; ws.row_dimensions[4].height = 16
    _w(ws, 1, 2, COMPANY.get("upper") or COMPANY["name"], bold=True, size=14, color=GOLD)
    _w(ws, 2, 2, COMPANY["address"], size=9, color="57534E")
    _w(ws, 3, 2, COMPANY["phone"], size=9, color="57534E")
    _w(ws, 4, 2, COMPANY["email"], size=9, color="57534E")
    ws.merge_cells("E1:F2"); _w(ws, 1, 5, title, bold=True, size=16, color=GOLD, align="right")
    r = 6
    for k, v in meta:
        _w(ws, r, 2, k, bold=True, color=GOLD, size=9); _w(ws, r, 3, v); ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6); r += 1
    r += 1
    heads = ["#", "İşin Cinsi / Açıklama", "Birim", "Miktar", "Birim Fiyat", "Toplam"]
    for i, h in enumerate(heads, 1):
        _w(ws, r, i, h, bold=True, fill="FAF5E8", border=True, align="center" if i != 2 else "left")
    r += 1
    for n, it in enumerate(items, 1):
        _w(ws, r, 1, n, border=True, align="center"); _w(ws, r, 2, it["description"], border=True)
        _w(ws, r, 3, it.get("unit", ""), border=True, align="center"); _w(ws, r, 4, it.get("quantity", 0), border=True, align="right", fmt="#,##0.##")
        _w(ws, r, 5, it.get("unit_price", 0), border=True, align="right", fmt=MONEY); _w(ws, r, 6, it.get("total", 0), border=True, align="right", fmt=MONEY)
        r += 1
    r += 1
    for k, v, bold in totals:
        _w(ws, r, 5, k, bold=bold, color=GOLD if bold else "1C1917", align="right"); _w(ws, r, 6, v, bold=bold, color=GOLD if bold else "1C1917", fmt=MONEY, align="right"); r += 1
    for head, text in sections:
        if not text: continue
        r += 1; _w(ws, r, 2, head, bold=True, color=GOLD, size=11); r += 1
        for line in str(text).split("\n"):
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
            c = _w(ws, r, 2, line, align="left"); c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = max(15, 15 * (1 + len(line) // 95))
            r += 1
    r += 3
    _w(ws, r, 2, signatures[0], bold=True, align="center"); ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
    c = _w(ws, r, 5, signatures[1], bold=True, align="center"); c.alignment = Alignment(wrap_text=True, horizontal="center"); ws.row_dimensions[r].height = 32
    ws.cell(row=r, column=2).border = Border(top=Side(style="thin", color="78716C"))
    for col in (5, 6): ws.cell(row=r, column=col).border = Border(top=Side(style="thin", color="78716C"))
    ws.page_setup.orientation = "portrait"; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True
    buf = BytesIO(); wb.save(buf)
    return buf.getvalue()
